import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from rw.continuation import continue_docx


class ContinuationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.folder = Path(self.temp.name)
        self.source = self.folder / 'latest.docx'; self.output = self.folder / 'next.docx'
        doc = Document(); p = doc.add_paragraph(); p.add_run('用户').bold = True; p.add_run('手工修改'); doc.add_table(rows=1, cols=1).cell(0, 0).text = '改善流程'; doc.save(self.source)
    def tearDown(self): self.temp.cleanup()
    def test_without_history_copies_latest_byte_for_byte(self):
        report = continue_docx(self.source, self.output)
        self.assertTrue(report['ok']); self.assertEqual(self.source.read_bytes(), self.output.read_bytes())
    def test_exact_replacement_preserves_styles_and_original(self):
        digest = hashlib.sha256(self.source.read_bytes()).hexdigest()
        report = continue_docx(self.source, self.output, {'用户手工修改': '新的手工修改', '改善流程': '优化流程'})
        self.assertTrue(report['ok']); self.assertEqual(Document(self.output).paragraphs[0].text, '新的手工修改')
        self.assertTrue(Document(self.output).paragraphs[0].runs[0].bold)
        self.assertEqual(Document(self.output).tables[0].cell(0, 0).text, '优化流程')
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), digest)
        self.assertEqual(len(report['changes']), 2); self.assertEqual(len(report['pending_facts']), 2)
    def test_missing_ambiguous_and_same_source_rejected_before_write(self):
        with self.assertRaises(ValueError): continue_docx(self.source, self.output, {'不存在': '新'})
        self.assertFalse(self.output.exists())
        with self.assertRaises(ValueError): continue_docx(self.source, self.source)
        doc = Document(self.source); doc.add_paragraph('改善流程'); doc.save(self.source)
        with self.assertRaises(ValueError): continue_docx(self.source, self.output, {'改善流程': '新'})
    def test_locations_approve_repeats_without_chaining_new_text(self):
        doc = Document(self.source); doc.add_paragraph('A B A'); doc.save(self.source)
        report = continue_docx(self.source, self.output, [{'old': 'A', 'new': 'B', 'location': 'body:2', 'count': 2}, {'old': 'B', 'new': 'C', 'location': 'body:2'}])
        self.assertTrue(report['ok']); self.assertEqual(Document(self.output).paragraphs[1].text, 'B C B')
    def test_existing_output_not_overwritten(self):
        self.output.write_bytes(b'important')
        with self.assertRaises(FileExistsError): continue_docx(self.source, self.output)
        self.assertEqual(self.output.read_bytes(), b'important')
    def test_multiline_replacement_round_trips_breaks(self):
        doc = Document(self.source); doc.add_paragraph('第一行\n第二行'); doc.save(self.source)
        report = continue_docx(self.source, self.output, {'第一行\n第二行': '新第一行\n新第二行'})
        self.assertTrue(report['ok']); self.assertEqual(Document(self.output).paragraphs[1].text, '新第一行\n新第二行')
        self.assertEqual(len(Document(self.output).paragraphs[1]._p.xpath('.//w:br')), 1)
    def test_body_edit_does_not_create_header_footer_parts_or_references(self):
        with zipfile.ZipFile(self.source) as package:
            before_parts = set(package.namelist())
            before_content = {name: package.read(name) for name in before_parts}
        before = Document(self.source)._element.xpath('.//w:headerReference | .//w:footerReference')
        report = continue_docx(self.source, self.output, {'用户手工修改': '仅修改正文'})
        self.assertTrue(report['ok'])
        with zipfile.ZipFile(self.output) as package:
            after_parts = set(package.namelist())
            after_content = {name: package.read(name) for name in after_parts}
        after = Document(self.output)._element.xpath('.//w:headerReference | .//w:footerReference')
        self.assertEqual(before_parts, after_parts)
        self.assertEqual(len(before), len(after))
        self.assertEqual(['word/document.xml'], sorted(name for name in before_parts if before_content[name] != after_content[name]))
    def test_existing_header_can_be_edited_without_new_variants(self):
        doc = Document(self.source); doc.sections[0].header.paragraphs[0].text = '现有页眉'; doc.save(self.source)
        with zipfile.ZipFile(self.source) as package: before_parts = set(package.namelist())
        report = continue_docx(self.source, self.output, [{'old': '现有页眉', 'new': '修改页眉', 'location': 'header:0:0'}])
        self.assertTrue(report['ok'])
        with zipfile.ZipFile(self.output) as package: self.assertEqual(before_parts, set(package.namelist()))
        self.assertEqual(Document(self.output).sections[0].header.paragraphs[0].text, '修改页眉')
    def test_linked_section_header_is_visited_once_without_new_reference(self):
        doc = Document(self.source); doc.sections[0].header.paragraphs[0].text = '共享页眉'
        doc.add_section(); doc.add_paragraph('第二节正文'); doc.save(self.source)
        before_refs = [node.xml for node in Document(self.source)._element.xpath('.//w:headerReference | .//w:footerReference')]
        report = continue_docx(self.source, self.output, {'共享页眉': '修改共享页眉'})
        self.assertTrue(report['ok'])
        self.assertEqual(1, len(report['changes']))
        after_doc = Document(self.output)
        self.assertEqual(before_refs, [node.xml for node in after_doc._element.xpath('.//w:headerReference | .//w:footerReference')])
        self.assertEqual('修改共享页眉', after_doc.sections[1].header.paragraphs[0].text)
