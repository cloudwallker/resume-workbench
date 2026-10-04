import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from unittest.mock import patch

from docx import Document
from rw import rendering


def sample(language='zh'):
    return {'profile': {'name': '虚构姓名' if language == 'zh' else 'Alex Example',
                        'email': 'example@example.com', 'headline': '运营经理',
                        'summary': '带领团队优化流程',
                        'links': [{'label': 'Portfolio', 'url': 'https://example.com/work'}]},
            'target': {'language': language, 'seniority': 'experienced', 'page_limit': 1},
            'sections': [{'id': 'work', 'title': '工作经历' if language == 'zh' else 'Experience',
                          'entries': [{'id': 'w1', 'title': '运营优化', 'organization': '示例公司',
                                       'role': '经理', 'start': '2020', 'end': '2024',
                                       'bullets': ['完成 12 个项目，效率提升 20%'],
                                       'url': 'https://example.com/project'}]}]}


class RenderingTests(unittest.TestCase):
    def test_pdf_font_spelling_resolves_to_installed_family(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'font-alias.docx'
            with patch.object(rendering, '_installed_font_names', return_value=['Microsoft YaHei'], create=True):
                result = rendering.render_docx(sample(), path, {'body_font': 'MicrosoftYaHei'})
            self.assertTrue(result['ok'], result)
            self.assertEqual('Microsoft YaHei', Document(path).styles['Normal'].font.name)
            self.assertTrue(result['adjustments'])
            self.assertEqual('MicrosoftYaHei', result['adjustments'][0]['from'])
    def test_editable_bilingual_body_and_clean_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            for language in ('zh', 'en'):
                path = Path(tmp) / (language + '.docx')
                result = rendering.render_docx(sample(language), path)
                self.assertTrue(result['ok'])
                doc = Document(path)
                text = '\n'.join(p.text for p in doc.paragraphs)
                self.assertIn(sample(language)['profile']['name'], text)
                self.assertIn('完成 12 个项目，效率提升 20%', text)
                self.assertEqual('', doc.core_properties.author)
                with ZipFile(path) as archive:
                    body = archive.read('word/document.xml').decode()
                    self.assertNotIn('<w:ins', body)
                    self.assertNotIn('<w:del', body)
                    self.assertNotIn('word/comments.xml', archive.namelist())
                targets = [r.target_ref for r in doc.part.rels.values() if r.is_external]
                self.assertIn('https://example.com/work', targets)

    def test_layout_is_applied_without_shrinking_for_page_limit(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'layout.docx'
            layout = {'paper_width_cm': 21.0, 'paper_height_cm': 29.7,
                      'margins_cm': {'top': 2, 'bottom': 2, 'left': 2.2, 'right': 2.2},
                      'body_font': 'Arial', 'body_size_pt': 12, 'heading_size_pt': 16,
                      'heading_color': '#123456'}
            rendering.render_docx(sample(), path, layout)
            doc = Document(path)
            self.assertAlmostEqual(21.0, doc.sections[0].page_width.cm, places=2)
            self.assertAlmostEqual(2.2, doc.sections[0].left_margin.cm, places=2)
            self.assertEqual(12, doc.styles['Normal'].font.size.pt)
            self.assertEqual('Arial', doc.styles['Normal'].font.name)

    def test_existing_output_is_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'original.docx'
            path.write_bytes(b'original')
            result = rendering.render_docx(sample(), path)
            self.assertFalse(result['ok'])
            self.assertEqual(b'original', path.read_bytes())

    def test_unknown_font_is_reported_without_silent_substitution(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'unknown.docx'
            result = rendering.render_docx(sample(), path, {'body_font': 'ResumeWorkbenchNonexistentFont'})
            self.assertTrue(result['ok'], result)
            self.assertIn('font_validation', result)
            self.assertNotEqual(True, result['font_validation']['available'])
            self.assertEqual('ResumeWorkbenchNonexistentFont', Document(path).styles['Normal'].font.name)
            self.assertTrue(any('字体' in item for item in result['warnings']))
