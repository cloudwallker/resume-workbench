import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from rw.quality import check_outputs


class QualityFieldTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.folder = Path(self.temporary.name)
        self.docx = self.folder / 'latest.docx'
        self.pdf = self.folder / 'final.pdf'

    def tearDown(self):
        self.temporary.cleanup()

    def write_pdf(self, pages, links=None):
        """Offline fixture with real extractable text and matching source hashes."""
        writer = PdfWriter()
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        font_reference = writer._add_object(font)
        for lines in pages:
            page = writer.add_blank_page(width=595, height=842)
            page[NameObject('/Resources')] = DictionaryObject({
                NameObject('/Font'): DictionaryObject({NameObject('/F1'): font_reference})})
            commands = ['BT /F1 11 Tf 50 780 Td']
            for index, line in enumerate(lines):
                if index:
                    commands.append('0 -30 Td')
                literal = line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
                commands.append('(' + literal + ') Tj')
            commands.append('ET')
            stream = DecodedStreamObject()
            stream.set_data(' '.join(commands).encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
        for page_number, url in links or []:
            writer.add_uri(page_number, url, [50, 50, 250, 70])
        with self.pdf.open('wb') as output:
            writer.write(output)
        self.pdf.with_suffix('.pdf.source.json').write_text(json.dumps({
            'source_sha256': hashlib.sha256(self.docx.read_bytes()).hexdigest(),
            'pdf_sha256': hashlib.sha256(self.pdf.read_bytes()).hexdigest(),
            'backend': 'offline-fixture'}), encoding='utf-8')

    def add_page_field(self, paragraph):
        for field_type, value in [('begin', None), (None, ' PAGE '),
                                  ('separate', None), (None, '1'), ('end', None)]:
            run = OxmlElement('w:r')
            if field_type:
                node = OxmlElement('w:fldChar')
                node.set(qn('w:fldCharType'), field_type)
            else:
                node = OxmlElement('w:instrText' if value == ' PAGE ' else 'w:t')
                node.text = value
            run.append(node)
            paragraph._p.append(run)

    def test_page_field_instruction_is_not_visible_resume_content(self):
        document = Document()
        document.add_paragraph('Latest resume')
        self.add_page_field(document.sections[0].footer.paragraphs[0])
        document.save(self.docx)
        self.write_pdf([['Latest resume', '1']])
        report = check_outputs(self.docx, self.pdf)
        self.assertTrue(report['ok'], report['issues'])
        self.assertEqual('Latest resume\n1', report['docx_text'])

    def test_body_field_instruction_is_not_a_tracked_insertion(self):
        document = Document()
        paragraph = document.add_paragraph('Document page ')
        self.add_page_field(paragraph)
        document.save(self.docx)
        self.write_pdf([['Document page 1']])
        report = check_outputs(self.docx, self.pdf)
        self.assertTrue(report['ok'], report['issues'])

    def test_multiple_pages_accept_interleaved_headers_and_footers(self):
        document = Document()
        document.add_paragraph('First page')
        document.add_page_break()
        document.add_paragraph('Second page')
        document.sections[0].header.paragraphs[0].text = 'Contact details'
        document.sections[0].footer.paragraphs[0].text = 'Footer'
        document.save(self.docx)
        self.write_pdf([['Contact details', 'First page', 'Footer'],
                        ['Contact details', 'Second page', 'Footer']])
        report = check_outputs(self.docx, self.pdf)
        self.assertTrue(report['ok'], report['issues'])

    def test_missing_body_paragraph_is_reported_without_payload(self):
        document = Document()
        document.add_paragraph('First page')
        document.add_paragraph('Missing work achievement')
        document.sections[0].footer.paragraphs[0].text = 'Footer'
        document.save(self.docx)
        self.write_pdf([['First page', 'Footer']])
        report = check_outputs(self.docx, self.pdf)
        self.assertFalse(report['ok'])
        missing = [item['value'] for item in report['issues'] if item['code'] == 'pdf_missing_content']
        self.assertEqual(['Missing work achievement'], missing)

    def test_visible_tabs_and_breaks_are_preserved_in_docx_report(self):
        document = Document()
        document.add_paragraph('Alpha\tBeta\nGamma')
        document.save(self.docx)
        self.write_pdf([['Alpha Beta', 'Gamma']])
        report = check_outputs(self.docx, self.pdf)
        self.assertTrue(report['ok'], report['issues'])
        self.assertEqual('Alpha\tBeta\nGamma', report['docx_text'])

    def test_payload_content_checks_remain_required(self):
        document = Document()
        document.add_paragraph('Latest resume')
        document.save(self.docx)
        self.write_pdf([['Latest resume']])
        report = check_outputs(self.docx, self.pdf, {'profile': {'name': 'Missing candidate'}})
        codes = {item['code'] for item in report['issues']}
        self.assertFalse(report['ok'])
        self.assertIn('missing_content', codes)
        self.assertIn('pdf_missing_content', codes)

    def test_payload_hyperlink_checks_remain_required(self):
        document = Document()
        document.add_paragraph('Latest resume')
        paragraph = document.add_paragraph()
        link = OxmlElement('w:hyperlink')
        url = 'https://fiction.invalid/portfolio'
        link.set(qn('r:id'), document.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Portfolio'
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document.save(self.docx)
        self.write_pdf([['Latest resume', 'Portfolio']])
        payload = {'profile': {'name': 'Latest resume', 'links': [{'label': 'Portfolio', 'url': url}]}}
        report = check_outputs(self.docx, self.pdf, payload)
        codes = {item['code'] for item in report['issues']}
        self.assertFalse(report['ok'])
        self.assertNotIn('missing_link', codes)
        self.assertIn('pdf_missing_link', codes)

    def test_existing_header_hyperlink_is_checked_in_its_story_part(self):
        document = Document()
        document.add_paragraph('Latest resume')
        paragraph = document.sections[0].header.paragraphs[0]
        url = 'https://fiction.invalid/portfolio'
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = 'Portfolio'
        run.append(text)
        link.append(run)
        paragraph._p.append(link)
        document.save(self.docx)
        self.write_pdf([['Portfolio', 'Latest resume']], [(0, url)])
        payload = {'profile': {'name': 'Latest resume', 'links': [{'label': 'Portfolio', 'url': url}]}}
        report = check_outputs(self.docx, self.pdf, payload)
        self.assertTrue(report['ok'], report['issues'])
        self.assertIn(url, report['docx_links'])


if __name__ == '__main__':
    unittest.main()
