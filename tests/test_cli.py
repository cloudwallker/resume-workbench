import json
from pathlib import Path
import subprocess
import sys
import tempfile
import os
import unittest

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / 'skills' / 'resume-workbench' / 'scripts' / 'resume.py'

def fictional_master():
    return {
        'schema_version': '1.0',
        'profile': {'name': '示例求职者', 'email': 'candidate@example.com', 'headline': '运营专员'},
        'target': {'role':'运营专员', 'language':'zh', 'market':'CN', 'seniority':'experienced', 'paper':'A4'},
        'sources': [{'id':'user-1', 'type':'user', 'reference':'虚构样例本人陈述'}],
        'experiences': [{
            'id':'work-1','kind':'work','title':'运营专员','organization':'示例服务机构','role':'负责内容整理',
            'start':'2023-06','end':'2025-08','summary':'','tags':['运营','内容'], 'status':'confirmed','source_ids':['user-1'],
            'bullets':[{'id':'b-1','text':'整理活动报名资料并交付周报，供团队复盘使用。','status':'confirmed',
                        'contribution':'individual','source_ids':['user-1']}]
        }]
    }

class CliTests(unittest.TestCase):
    @unittest.skipUnless(os.name == 'nt', 'Windows ACL behavior')
    def test_output_directory_inherits_workspace_access_rules(self):
        with tempfile.TemporaryDirectory() as temp:
            master=Path(temp)/'master.json'; master.write_text(json.dumps(fictional_master(),ensure_ascii=False),encoding='utf-8')
            result=self.call('build','--master',master,'--outdir',Path(temp)/'output','--docx-only')
            self.assertEqual(0,result.returncode,result.stdout+result.stderr)
            report=json.loads(result.stdout)
            access=subprocess.run(['icacls.exe',report['directory']],capture_output=True)
            self.assertEqual(0,access.returncode)
            self.assertIn(b'(I)',access.stdout,'Output needs inherited workspace permissions for desktop Word and user access.')

    def test_json_report_is_utf8_even_in_an_ascii_console(self):
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'input.txt'; source.write_text('中文简历 😀',encoding='utf-8')
            environment=dict(os.environ, PYTHONIOENCODING='ascii')
            result=subprocess.run([sys.executable,str(CLI),'import',str(source)],capture_output=True,env=environment)
            self.assertEqual(0,result.returncode,result.stderr.decode('utf-8',errors='replace'))
            self.assertEqual('中文简历 😀',json.loads(result.stdout.decode('utf-8'))['text'])
    def call(self, *args):
        return subprocess.run([sys.executable, '-X', 'utf8', str(CLI), *map(str,args)], capture_output=True, text=True, encoding='utf-8')

    def test_doctor_returns_json_capability_report(self):
        result=self.call('doctor')
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertIsInstance(json.loads(result.stdout),dict)

    def test_invalid_master_is_reported_without_output(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'bad.json'
            path.write_text(json.dumps({'schema_version':'1.0','profile':{},'experiences':[]}),encoding='utf-8')
            result=self.call('build','--master',path,'--outdir',Path(temp)/'out','--docx-only')
            self.assertNotEqual(result.returncode,0)
            self.assertTrue(result.stdout.strip().startswith('{'),result.stderr)
            self.assertIn('error',json.loads(result.stdout))
            self.assertFalse(list(Path(temp).rglob('*.docx')))

    def test_docx_only_is_explicit_and_preserves_previous_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'master.json'
            path.write_text(json.dumps(fictional_master(),ensure_ascii=False),encoding='utf-8')
            first=self.call('build','--master',path,'--outdir',Path(temp)/'out','--docx-only')
            self.assertEqual(first.returncode,0,first.stderr)
            a=json.loads(first.stdout)
            self.assertEqual(a['status'],'docx_only')
            original=Path(a['docx']).read_bytes()
            second=self.call('build','--master',path,'--outdir',Path(temp)/'out','--docx-only')
            self.assertEqual(second.returncode,0,second.stderr)
            b=json.loads(second.stdout)
            self.assertNotEqual(a['docx'],b['docx'])
            self.assertEqual(Path(a['docx']).read_bytes(),original)

    def test_conversion_unavailable_is_partial_not_complete(self):
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'master.json'
            path.write_text(json.dumps(fictional_master(),ensure_ascii=False),encoding='utf-8')
            result=self.call('build','--master',path,'--outdir',Path(temp)/'out','--backend','none')
            self.assertEqual(result.returncode,2,result.stderr)
            self.assertTrue(result.stdout.strip().startswith('{'),result.stderr)
            report=json.loads(result.stdout)
            self.assertEqual(report['status'],'partial')
            self.assertTrue(Path(report['docx']).is_file())
            self.assertFalse(report.get('pdf'))
            self.assertFalse(report['conversion']['ok'])

    def test_import_template_does_not_assert_confirmed_facts(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp)/'模板.docx'
            document=Document(); document.add_paragraph('模板示例姓名'); document.add_paragraph('工作经历')
            document.save(path)
            result=self.call('import',path,'--purpose','template')
            self.assertEqual(result.returncode,0,result.stderr)
            report=json.loads(result.stdout)
            self.assertEqual(report['purpose'],'template')
            self.assertNotIn('confirmed',json.dumps(report,ensure_ascii=False))

    def test_continue_without_history_preserves_original(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'最新简历.docx'
            doc=Document(); p=doc.add_paragraph(); r=p.add_run('旧文案'); r.bold=True; doc.save(source)
            original=source.read_bytes()
            changes=Path(temp)/'changes.json'
            changes.write_text(json.dumps([{'old':'旧文案','new':'新文案'}],ensure_ascii=False),encoding='utf-8')
            result=self.call('continue',source,'--outdir',Path(temp)/'out','--replacements',changes,'--docx-only')
            self.assertEqual(result.returncode,0,result.stderr)
            report=json.loads(result.stdout)
            self.assertEqual(source.read_bytes(),original)
            modified=Document(report['docx'])
            self.assertEqual(modified.paragraphs[0].text,'新文案')
            self.assertTrue(modified.paragraphs[0].runs[0].bold)

    def test_failed_inspection_has_nonzero_exit_and_json(self):
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'source.bin'; source.write_bytes(b'unsupported')
            for command in ('import','analyze-template'):
                result=self.call(command,source)
                self.assertEqual(result.returncode,2,result.stdout+result.stderr)
                self.assertFalse(json.loads(result.stdout)['ok'])

    def test_broken_docx_returns_safe_json_without_traceback(self):
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'broken.docx'; source.write_bytes(b'not a docx')
            result=self.call('import',source)
            self.assertEqual(result.returncode,2,result.stderr)
            self.assertIn('error',json.loads(result.stdout))
            self.assertNotIn('Traceback',result.stderr)

    def test_export_existing_word_retains_bytes_and_reports_missing_converter(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as temp:
            source=Path(temp)/'final.docx'; doc=Document(); doc.add_paragraph('最新版'); doc.save(source)
            original=source.read_bytes()
            result=self.call('export','--docx',source,'--outdir',Path(temp)/'preview','--backend','none')
            self.assertEqual(result.returncode,2,result.stdout+result.stderr)
            report=json.loads(result.stdout)
            self.assertEqual('partial',report['status'])
            self.assertEqual('export',report['operation'])
            self.assertEqual(original,source.read_bytes())
            self.assertEqual(original,Path(report['docx']).read_bytes())

if __name__=='__main__':
    unittest.main()
