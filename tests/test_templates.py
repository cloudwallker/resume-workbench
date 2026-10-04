import tempfile
import unittest
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from pypdf import PdfWriter
from rw.templates import analyze_template, adapt_docx, pdf_layout_profile

PAYLOAD = {'profile': {'name': '李明', 'phone': '123456', 'email': 'fiction@example.invalid'}, 'target': {'language': 'zh'}, 'sections': [{'id': 'work', 'title': '工作经历', 'entries': [{'id': 'e1', 'title': '运营经理', 'organization': '虚构公司', 'role': '', 'start': '2022', 'end': '2025', 'summary': '', 'bullets': ['改善运营流程']}]}]}


def plain_pdf(path, content):
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter(); page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(content); page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as output: writer.write(output)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.folder = Path(self.temp.name)
    def tearDown(self): self.temp.cleanup()
    def test_placeholder_split_across_runs_keeps_style(self):
        path = self.folder / 'template.docx'; output = self.folder / 'new.docx'
        doc = Document(); p = doc.add_paragraph(); p.add_run('{{profile.').bold = True; p.add_run('name}}'); doc.add_paragraph('{{sections}}'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertTrue(report['ok'])
        actual = Document(output)
        self.assertEqual(actual.paragraphs[0].text, '李明')
        self.assertTrue(actual.paragraphs[0].runs[0].bold)
        self.assertIn('改善运营流程', '\n'.join(p.text for p in actual.paragraphs))
    def test_heading_zone_and_simple_label_table_remove_samples(self):
        path = self.folder / 'template.docx'; output = self.folder / 'new.docx'
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '姓名'; table.cell(0, 1).text = '样例姓名'
        table.cell(1, 0).text = '电话'; table.cell(1, 1).text = '样例号码'
        doc.add_paragraph('工作经历', style='Heading 1'); doc.add_paragraph('样例公司'); doc.add_paragraph('样例职责'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertTrue(report['ok'], report)
        result = Document(output)
        text = '\n'.join(p.text for p in result.paragraphs) + str([[c.text for c in r.cells] for r in result.tables[0].rows])
        self.assertNotIn('样例', text); self.assertIn('虚构公司', text); self.assertEqual(result.tables[0].cell(0, 1).text, '李明')
    def test_unmapped_content_blocks_output_and_accepts_explicit_mapping(self):
        path = self.folder / 'template.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('某人简历'); doc.add_paragraph('样例公司'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertFalse(report['ok']); self.assertTrue(report['mapping_required']); self.assertFalse(output.exists())
        mapped = adapt_docx(path, PAYLOAD, output, {'paragraphs': {'0': 'profile.name', '1': 'section:work'}})
        self.assertTrue(mapped['ok']); self.assertNotIn('样例公司', '\n'.join(p.text for p in Document(output).paragraphs))
    def test_complex_columns_warn_and_do_not_silently_adapt(self):
        path = self.folder / 'columns.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('{{profile.name}}')
        doc.sections[0]._sectPr.xpath('./w:cols')[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '2'); doc.save(path)
        report = analyze_template(path)
        self.assertIn('columns', report['issues']); self.assertFalse(adapt_docx(path, PAYLOAD, output)['ok'])
    def test_same_path_rejected(self):
        path = self.folder / 'template.docx'; Document().save(path)
        with self.assertRaises(ValueError): adapt_docx(path, PAYLOAD, path)
    def test_pdf_reads_real_paper_and_scan_status(self):
        path = self.folder / 'scan.pdf'; writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
        with path.open('wb') as stream: writer.write(stream)
        report = pdf_layout_profile(path)
        self.assertAlmostEqual(report['layout']['paper_width_cm'], 21.59, places=2)
        self.assertTrue(report['requires_ocr']); self.assertFalse(report['ok'])
    def test_pdf_text_font_bounds_and_heading_features(self):
        path = self.folder / 'text.pdf'
        writer = PdfWriter(); page = writer.add_blank_page(width=600, height=800)
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 20 Tf 0.2 0.4 0.6 rg 50 750 Td (EXPERIENCE) Tj /F1 11 Tf 0 0 0 rg 0 -40 Td (Operations and service delivery) Tj 0 -20 Td (Managed a fictional team) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        with path.open('wb') as output: writer.write(output)
        report = pdf_layout_profile(path)
        self.assertTrue(report['ok'], report)
        self.assertEqual(report['layout']['body_font'], 'Helvetica'); self.assertEqual(report['layout']['body_size_pt'], 11)
        self.assertEqual(report['layout']['heading_size_pt'], 20); self.assertEqual(report['layout']['heading_color'], '336699')
        self.assertAlmostEqual(report['layout']['margins_cm']['left'], 50 * 2.54 / 72, places=2)
    def test_pdf_section_heading_style_wins_over_larger_name(self):
        path = self.folder / 'section-style.pdf'; writer = PdfWriter(); page = writer.add_blank_page(width=600, height=800)
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject(); stream.set_data(b'BT /F1 18 Tf 0 0 0 rg 50 750 Td (FICTIONAL NAME) Tj /F1 12 Tf 0.2 0.4 0.6 rg 0 -40 Td (EXPERIENCE) Tj /F1 10.5 Tf 0 0 0 rg 0 -25 Td (Operations and service delivery) Tj 0 -20 Td (Managed a fictional team and service process) Tj /F1 12 Tf 0.2 0.4 0.6 rg 0 -40 Td (EDUCATION) Tj /F1 10.5 Tf 0 0 0 rg 0 -25 Td (Business administration degree at fictional school) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        with path.open('wb') as output: writer.write(output)
        report = pdf_layout_profile(path)
        self.assertEqual(10.5, report['layout']['body_size_pt'])
        self.assertEqual(12, report['layout']['heading_size_pt'])
        self.assertEqual('336699', report['layout']['heading_color'])
        self.assertEqual(['work', 'education'], [heading['section'] for heading in report['headings']])
    def test_short_single_column_gaps_are_evidence_not_all_physical_margins(self):
        path = self.folder / 'short.pdf'
        plain_pdf(path, b'BT /F1 10.5 Tf 50 750 Td (Fictional operations resume) Tj 0 -25 Td (Short service experience) Tj ET')
        report = pdf_layout_profile(path); layout = report['layout']
        self.assertEqual(layout['margins_cm']['left'], layout['margins_cm']['right'])
        gaps = layout['measured_content_gaps_cm']
        self.assertGreater(gaps['right'], 2 * gaps['left']); self.assertGreater(gaps['bottom'], 2 * gaps['top'])
        self.assertEqual(gaps['left'], layout['margins_cm']['right'])
        self.assertEqual(gaps['top'], layout['margins_cm']['bottom'])
        self.assertEqual(report['pages'][0]['bounds_pt'], layout['text_bounds_pt'][0]['bounds_pt'])
        self.assertEqual(10.5, layout['body_size_pt'])
        self.assertEqual({'margins_cm.right', 'margins_cm.bottom'}, {change['field'] for change in report['adjustments']})
        self.assertTrue(any('对称' in warning for warning in report['warnings']))
    def test_multicolumn_gaps_are_retained_and_uncertainty_reported(self):
        path = self.folder / 'two-columns.pdf'
        plain_pdf(path, b'BT /F1 10.5 Tf 50 750 Td (Left) Tj 300 0 Td (Right) Tj -300 -25 Td (Left) Tj 300 0 Td (Right) Tj -300 -25 Td (Left) Tj 300 0 Td (Right) Tj ET')
        report = pdf_layout_profile(path)
        self.assertIn('columns', report['issues'])
        self.assertEqual(report['layout']['measured_content_gaps_cm'], report['layout']['margins_cm'])
        self.assertEqual([], report['adjustments'])
        self.assertTrue(any('多栏' in warning for warning in report['warnings']))
    def test_vector_layout_gaps_are_not_reinterpreted_as_symmetry(self):
        path = self.folder / 'vector.pdf'
        plain_pdf(path, b'50 50 400 700 re S BT /F1 10.5 Tf 50 750 Td (Short fictional text) Tj ET')
        report = pdf_layout_profile(path)
        self.assertEqual(report['layout']['measured_content_gaps_cm'], report['layout']['margins_cm'])
        self.assertEqual([], report['adjustments'])
        self.assertTrue(any('推断' in warning for warning in report['warnings']))
    def test_explicit_header_mapping_and_template_author_are_cleared(self):
        path = self.folder / 'header.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('{{profile.name}}'); doc.sections[0].header.paragraphs[0].text = '样例页眉'
        doc.core_properties.author = '样例作者'; doc.save(path)
        self.assertFalse(adapt_docx(path, PAYLOAD, output)['ok'])
        report = adapt_docx(path, PAYLOAD, output, {'locations': {'header:0:0': 'profile.contact'}})
        self.assertTrue(report['ok'], report)
        actual = Document(output)
        self.assertNotIn('样例', actual.sections[0].header.paragraphs[0].text)
        self.assertEqual(actual.core_properties.author, '')
    def test_unknown_heading_is_boundary_and_requires_mapping(self):
        path = self.folder / 'ambiguous.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('工作经历', 'Heading 1'); doc.add_paragraph('工作样例')
        doc.add_paragraph('未知栏目', 'Heading 1'); doc.add_paragraph('不能猜测的样例'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertFalse(report['ok']); self.assertFalse(output.exists())
    def test_placeholder_inside_heading_zone_not_removed_as_sample(self):
        path = self.folder / 'heading.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('工作经历', 'Heading 1'); doc.add_paragraph('{{section:work}}')
        doc.add_paragraph('{{profile.email}}'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertTrue(report['ok'], report)
        self.assertIn('fiction@example.invalid', '\n'.join(p.text for p in Document(output).paragraphs))
    def test_filled_hyperlink_does_not_keep_sample_url(self):
        path = self.folder / 'link.docx'; output = self.folder / 'new.docx'
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml.ns import qn
        doc = Document(); p = doc.add_paragraph(); link = OxmlElement('w:hyperlink')
        rid = doc.part.relate_to('https://sample.invalid/old', RT.HYPERLINK, is_external=True)
        link.set(qn('r:id'), rid); run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '{{profile.name}}'; run.append(text); link.append(run); p._p.append(link); doc.save(path)
        self.assertTrue(adapt_docx(path, PAYLOAD, output)['ok'])
        actual = Document(output)
        self.assertEqual(actual.paragraphs[0].text, '李明')
        self.assertNotIn('https://sample.invalid/old', [rel.target_ref for rel in actual.part.rels.values()])
    def test_payload_links_are_clickable_or_report_missing_slot(self):
        from copy import deepcopy
        payload = deepcopy(PAYLOAD); payload['profile']['links'] = [{'label': 'Portfolio', 'url': 'https://fiction.invalid/portfolio'}]
        payload['sections'][0]['entries'][0]['url'] = 'https://fiction.invalid/project'
        path = self.folder / 'links.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('{{profile.name}}'); doc.add_paragraph('{{sections}}'); doc.save(path)
        report = adapt_docx(path, payload, output)
        self.assertFalse(report['ok']); self.assertTrue(report['mapping_required'])
        doc.add_paragraph('{{profile.links}}'); doc.save(path)
        report = adapt_docx(path, payload, output)
        self.assertTrue(report['ok'], report)
        targets = [rel.target_ref for rel in Document(output).part.rels.values()]
        self.assertIn('https://fiction.invalid/portfolio', targets); self.assertIn('https://fiction.invalid/project', targets)
        self.assertTrue(Document(output)._element.xpath('.//w:br'))
    def test_each_entry_url_has_a_relationship(self):
        from copy import deepcopy
        payload = deepcopy(PAYLOAD)
        payload['sections'][0]['entries'][0]['url'] = 'https://fiction.invalid/one'
        payload['sections'][0]['entries'].append({'title': '第二项目', 'url': 'https://fiction.invalid/two', 'bullets': []})
        path = self.folder / 'links.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('{{sections}}'); doc.save(path)
        self.assertTrue(adapt_docx(path, payload, output)['ok'])
        targets = [rel.target_ref for rel in Document(output).part.rels.values()]
        self.assertIn('https://fiction.invalid/one', targets); self.assertIn('https://fiction.invalid/two', targets)
    def test_first_and_even_page_samples_require_explicit_mapping(self):
        path = self.folder / 'special-headers.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('{{sections}}')
        doc.sections[0].different_first_page_header_footer = True
        doc.sections[0].first_page_header.paragraphs[0].text = '首页样例姓名'
        doc.sections[0].even_page_footer.paragraphs[0].text = '偶数页样例姓名'; doc.save(path)
        self.assertFalse(adapt_docx(path, PAYLOAD, output)['ok'])
        report = adapt_docx(path, PAYLOAD, output, {'locations': {'header:first:0:0': 'profile.name', 'footer:even:0:0': 'profile.contact'}})
        self.assertTrue(report['ok'], report)
    def test_plain_label_paragraph_replaces_sample_value(self):
        path = self.folder / 'labels.docx'; output = self.folder / 'new.docx'
        doc = Document(); doc.add_paragraph('姓名：样例姓名'); doc.add_paragraph('联系方式：样例邮箱')
        doc.add_paragraph('所在地：样例城市'); doc.add_paragraph('个人简介：样例简介'); doc.add_paragraph('{{sections}}'); doc.save(path)
        report = adapt_docx(path, PAYLOAD, output)
        self.assertTrue(report['ok'], report)
        text = '\n'.join(p.text for p in Document(output).paragraphs)
        self.assertIn('姓名：李明', text); self.assertIn('fiction@example.invalid', text); self.assertNotIn('样例', text)
    def test_analyze_lists_only_existing_headers_and_footers(self):
        path = self.folder / 'no-header.docx'; doc = Document(); doc.add_paragraph('{{sections}}'); doc.save(path)
        report = analyze_template(path)
        self.assertEqual(['body:0'], [p['location'] for p in report['locations']])
        doc.sections[0].header.paragraphs[0].text = '{{profile.name}}'; doc.save(path)
        report = analyze_template(path)
        self.assertEqual(['body:0', 'header:0:0'], [p['location'] for p in report['locations']])
