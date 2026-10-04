import tempfile
import unittest
from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from rw.importing import inspect_input


class ImportingTests(unittest.TestCase):
    def test_docx_text_table_and_template_are_not_confirmed_facts(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / '中文模板.docx'
            doc = Document(); doc.add_paragraph('样例姓名'); doc.add_table(rows=1, cols=2).cell(0, 1).text = '样例公司'
            doc.save(path)
            report = inspect_input(path, purpose='template')
            self.assertIn('样例公司', report['text'])
            self.assertFalse(report['allow_fact_import'])
            self.assertEqual(report['facts'], [])
            self.assertEqual(len(report['structure']['tables']), 1)
            self.assertEqual(inspect_input(path)['source_status'], 'pending')

    def test_scanned_pdf_reports_ocr(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'scan.pdf'
            writer = PdfWriter(); writer.add_blank_page(width=612, height=792)
            with path.open('wb') as stream: writer.write(stream)
            report = inspect_input(path)
            self.assertTrue(report['requires_ocr'])
            self.assertFalse(report['ok'])

    def test_legacy_doc_and_invalid_purpose_are_explicit(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'old.doc'; path.write_bytes(b'old document')
            self.assertTrue(inspect_input(path)['requires_conversion'])
            with self.assertRaises(ValueError): inspect_input(path, 'anything')
    def test_reading_does_not_invent_blank_header_footer_locations(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'plain.docx'; doc = Document(); doc.add_paragraph('正文'); doc.save(path)
            report = inspect_input(path)
            self.assertEqual(['body:0'], [p['location'] for p in report['structure']['paragraphs']])
