import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from rw import quality, rendering
from test_rendering import sample


class QualityTests(unittest.TestCase):
    def test_placeholders_and_missing_numbers_and_links_are_failures(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'bad.docx'
            doc = Document()
            doc.add_paragraph('{{姓名}} 完成 10 个项目')
            doc.save(path)
            report = quality.check_outputs(path, payload=sample())
            self.assertFalse(report['ok'])
            codes = {x['code'] for x in report['issues']}
            self.assertIn('placeholder', codes)
            self.assertIn('missing_content', codes)
            self.assertIn('missing_number', codes)
            self.assertIn('missing_link', codes)

    def test_docx_only_has_no_false_pdf_delivery_claim(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'good.docx'
            rendering.render_docx(sample(), path)
            report = quality.check_outputs(path, payload=sample())
            self.assertTrue(report['ok'], report)
            self.assertFalse(report['dual_delivery'])
            self.assertTrue(report['visual_review_required'])

    def test_blank_pdf_and_exceeded_page_limit_detected_and_rendered(self):
        with tempfile.TemporaryDirectory() as tmp:
            path, pdf = Path(tmp) / 'good.docx', Path(tmp) / 'bad.pdf'
            rendering.render_docx(sample(), path)
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            writer.add_blank_page(width=595, height=842)
            with pdf.open('wb') as stream:
                writer.write(stream)
            report = quality.check_outputs(path, pdf, sample(), Path(tmp) / 'pages')
            codes = {x['code'] for x in report['issues']}
            self.assertIn('pdf_no_text', codes)
            self.assertIn('page_limit', codes)
            self.assertEqual(2, len(report['rendered_pages']))
            self.assertTrue(report['visual_review_required'])

    def test_changed_docx_marks_bound_pdf_stale(self):
        with tempfile.TemporaryDirectory() as tmp:
            path, pdf = Path(tmp) / 'good.docx', Path(tmp) / 'good.pdf'
            rendering.render_docx(sample(), path)
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with pdf.open('wb') as stream:
                writer.write(stream)
            pdf.with_suffix('.pdf.source.json').write_text(json.dumps({'source_sha256': 'old', 'pdf_sha256': 'old'}))
            report = quality.check_outputs(path, pdf)
            self.assertIn('stale_pdf', {x['code'] for x in report['issues']})
