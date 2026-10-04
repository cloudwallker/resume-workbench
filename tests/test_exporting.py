import tempfile
import unittest
import shutil
import subprocess
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from rw import exporting


class ExportingTests(unittest.TestCase):
    @unittest.skipUnless(shutil.which('powershell.exe'), 'Windows PowerShell 5 required')
    def test_word_cleanup_ref_signatures_work_in_windows_powershell(self):
        # A managed signature double reproduces COM's ref-object binder without
        # starting Word. Open/export succeed; cleanup must accept ref arguments.
        types = r'''
Add-Type -TypeDefinition @'
using System;
using System.IO;
public class RWCompatDocument {
    public void ExportAsFixedFormat(string output, int format) { File.WriteAllText(output, "pdf"); }
    public void SaveAs2(object output, object format) { File.WriteAllText((string)output, "docx"); }
    public void Close(ref object save) { RWCompatWord.Closed = true; }
}
public class RWCompatDocuments {
    public RWCompatDocument Open(object input, object confirm, object readOnly, object recent) {
        if (!(bool)readOnly || (bool)recent) throw new Exception("Unsafe open options");
        return new RWCompatDocument();
    }
}
public class RWCompatWord {
    public static bool Closed = false;
    public static bool QuitCalled = false;
    public bool Visible { get; set; }
    public int DisplayAlerts { get; set; }
    public int AutomationSecurity { get; set; }
    public RWCompatDocuments Documents = new RWCompatDocuments();
    public void Quit(ref object save) {
        if (!Closed || Visible || AutomationSecurity != 3) throw new Exception("Unsafe cleanup");
        QuitCalled = true;
    }
}
'@
'''
        script = exporting.WORD_SCRIPT.replace('New-Object -ComObject Word.Application', 'New-Object RWCompatWord')
        script = script.replace('[void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($document)', '$null = $document')
        script = script.replace('[void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($word)', '$null = $word')
        # Keep param first, append the type definitions after the param block.
        script = script.replace("$ErrorActionPreference = 'Stop'", "$ErrorActionPreference = 'Stop'\n" + types)
        script += '\nif (-not [RWCompatWord]::QuitCalled) { throw "Quit not called" }\n'
        with tempfile.TemporaryDirectory() as tmp:
            ps = Path(tmp) / 'compat.ps1'
            ps.write_text(script, encoding='utf-8-sig')
            for kind in ('pdf', 'docx'):
                output = Path(tmp) / ('姓名 & output.' + kind)
                result = subprocess.run([shutil.which('powershell.exe'), '-NoProfile', '-NonInteractive',
                                         '-File', str(ps), '-InputPath', "姓名 & $(bad) 'source.docx",
                                         '-OutputPath', str(output), '-Format', kind], capture_output=True,
                                        text=True, errors='replace', timeout=30, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
                self.assertEqual(0, result.returncode, result.stderr)
                self.assertEqual(kind, output.read_text())

    def test_missing_backend_preserves_docx_and_reports_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / '中文 & special.docx'
            Document().save(source)
            before = source.read_bytes()
            result = exporting.convert_to_pdf(source, Path(tmp) / 'result.pdf', backend='unsupported')
            self.assertFalse(result['ok'])
            self.assertEqual(before, source.read_bytes())
            self.assertFalse((Path(tmp) / 'result.pdf').exists())
            self.assertIn('recovery', result)

    def test_existing_pdf_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / 'source.docx', Path(tmp) / 'source.pdf'
            Document().save(source)
            target.write_bytes(b'original')
            self.assertFalse(exporting.convert_to_pdf(source, target)['ok'])
            self.assertEqual(b'original', target.read_bytes())

    def test_word_uses_argument_paths_and_hidden_owned_instance(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "姓名 & $(bad) '简历.docx"
            Document().save(source)
            target = Path(tmp) / 'out.pdf'
            observed = {}
            def run(command, **kwargs):
                observed['command'] = command
                script = Path(command[command.index('-File') + 1]).read_text(encoding='utf-8-sig')
                observed['script'] = script
                raise OSError('test conversion unavailable')
            with patch.object(exporting.subprocess, 'run', side_effect=run):
                result = exporting.convert_to_pdf(source, target, backend='word')
            self.assertFalse(result['ok'])
            self.assertIn(str(source.resolve()), observed['command'])
            self.assertNotIn(str(source.resolve()), observed['script'])
            self.assertIn('AutomationSecurity = 3', observed['script'])
            self.assertIn('$word.Visible = $false', observed['script'])
            self.assertNotIn('Stop-Process', observed['script'])

    def test_doctor_has_explicit_capability_report(self):
        report = exporting.doctor()
        self.assertIn('backends', report)
        self.assertIn('word', report['backends'])
        self.assertIn('libreoffice', report['backends'])

    def test_success_is_bound_to_source_and_valid_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / 'source.docx', Path(tmp) / 'final.pdf'
            Document().save(source)
            def convert(source, output, kind, temp):
                writer = PdfWriter()
                writer.add_blank_page(width=595, height=842)
                with output.open('wb') as stream:
                    writer.write(stream)
            with patch.object(exporting, '_word', side_effect=convert):
                result = exporting.convert_to_pdf(source, target, backend='word')
            self.assertTrue(result['ok'], result)
            self.assertEqual(result['source_sha256'], exporting._hash(source))
            self.assertTrue(target.with_suffix('.pdf.source.json').is_file())

    def test_invalid_pdf_bytes_not_delivered(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / 'source.docx', Path(tmp) / 'final.pdf'
            Document().save(source)
            def convert(source, output, kind, temp):
                output.write_bytes(b'not a pdf')
            with self.assertNoLogs('pypdf', level='WARNING'):
                with patch.object(exporting, '_word', side_effect=convert):
                    result = exporting.convert_to_pdf(source, target, backend='word')
            self.assertFalse(result['ok'])
            self.assertFalse(target.exists())

    def test_source_modified_during_conversion_not_delivered(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, target = Path(tmp) / 'source.docx', Path(tmp) / 'final.pdf'
            Document().save(source)
            def convert(source, output, kind, temp):
                output.write_bytes(b'candidate')
                document = Document(source)
                document.add_paragraph('modified while exporting')
                document.save(source)
            with patch.object(exporting, '_word', side_effect=convert):
                result = exporting.convert_to_pdf(source, target, backend='word')
            self.assertFalse(result['ok'])
            self.assertFalse(target.exists())
