"""Editable Word rendering. Page preferences never trigger silent font reduction."""
from datetime import datetime
import hashlib
import os
from pathlib import Path
import re
from urllib.parse import urlparse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def _font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)


def _link(paragraph, label, url):
    parsed = urlparse(str(url))
    if parsed.scheme not in ('https', 'http', 'mailto') or (parsed.scheme != 'mailto' and not parsed.netloc):
        raise ValueError('链接必须使用有效 http、https 或 mailto 地址')
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), paragraph.part.relate_to(str(url), RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '245B85'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = str(label or url); run.append(text)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def _clean_metadata(document):
    core = document.core_properties
    for key in ('author', 'last_modified_by', 'title', 'subject', 'keywords', 'comments',
                'category', 'content_status', 'identifier', 'language', 'version'):
        setattr(core, key, '')
    core.created = core.modified = datetime(2000, 1, 1)
    core.revision = 1


def _installed_font_names():
    """Read installed family names without loading or installing fonts."""
    if os.name != 'nt':
        return None
    import winreg
    names_found = set()
    readable = False
    for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            with winreg.OpenKey(hive, r'SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts') as key:
                readable = True
                index = 0
                while True:
                    try:
                        family, _, _ = winreg.EnumValue(key, index)
                    except OSError:
                        break
                    names = family.split(' (')[0].split(' & ')
                    names_found.update(names)
                    index += 1
        except OSError:
            continue
    return sorted(names_found) if readable else None


def _font_available(name):
    names = _installed_font_names()
    return any(n.casefold() == name.casefold() for n in names) if names is not None else None


def _resolve_font(name):
    names = _installed_font_names()
    if names is None:
        return name
    normalized = re.sub(r'\s+', '', name).casefold()
    matches = [n for n in names if re.sub(r'\s+', '', n).casefold() == normalized]
    return matches[0] if len(matches) == 1 else name


def render_docx(payload, output_path, layout=None):
    """Return a report and create a new DOCX, never replace an existing file."""
    output = Path(output_path).resolve()
    if output.exists():
        return {'ok': False, 'error': '输出文件已存在；请使用新文件名', 'docx_path': str(output)}
    try:
        language = payload.get('target', {}).get('language', 'zh')
        paper = payload.get('target', {}).get('paper', 'A4')
        defaults = {'paper_width_cm': 21.59 if paper == 'Letter' else 21,
                    'paper_height_cm': 27.94 if paper == 'Letter' else 29.7,
                    'margins_cm': dict(top=1.6, bottom=1.6, left=1.8, right=1.8),
                    'body_font': 'Microsoft YaHei' if language == 'zh' else 'Calibri',
                    'body_size_pt': 10.5 if language == 'zh' else 11,
                    'heading_size_pt': 12, 'heading_color': '245B85'}
        settings = {**defaults, **(layout or {})}
        requested_font = settings['body_font']
        settings['body_font'] = _resolve_font(requested_font)
        adjustments = list(settings.get('adjustments', []))
        if requested_font != settings['body_font']:
            adjustments.append({'field': 'body_font', 'from': requested_font, 'to': settings['body_font'],
                                'reason': 'Matched the installed family after normalizing whitespace in the PDF font name.'})
        settings['margins_cm'] = {**defaults['margins_cm'], **settings.get('margins_cm', {})}
        for key in ('paper_width_cm', 'paper_height_cm', 'body_size_pt', 'heading_size_pt'):
            if not isinstance(settings[key], (int, float)) or settings[key] <= 0:
                raise ValueError('版式尺寸和字号必须为正数')
        margins = settings['margins_cm']
        if any(not isinstance(v, (int, float)) or v < 0 for v in margins.values()):
            raise ValueError('页边距必须为非负数')
        if margins['left'] + margins['right'] >= settings['paper_width_cm'] or margins['top'] + margins['bottom'] >= settings['paper_height_cm']:
            raise ValueError('页边距不能占满纸张')
        color = str(settings['heading_color']).lstrip('#')
        if not re.fullmatch('[0-9a-fA-F]{6}', color):
            raise ValueError('标题颜色须为六位十六进制颜色')
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Cm(settings['paper_width_cm']), Cm(settings['paper_height_cm'])
        for key in ('top', 'bottom', 'left', 'right'):
            setattr(section, key + '_margin', Cm(margins[key]))
        normal = document.styles['Normal']
        _font(normal, settings['body_font'], settings['body_size_pt'])
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.08
        heading = document.styles['Heading 1']
        _font(heading, settings['body_font'], settings['heading_size_pt'])
        heading.font.color.rgb = RGBColor.from_string(color)
        heading.paragraph_format.space_before = Pt(9)
        heading.paragraph_format.space_after = Pt(4)
        heading.paragraph_format.keep_with_next = True
        _font(document.styles['List Bullet'], settings['body_font'], settings['body_size_pt'])
        profile = payload.get('profile', {})
        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(3)
        run = title.add_run(str(profile.get('name', '')))
        run.bold, run.font.size = True, Pt(max(18, settings['heading_size_pt'] + 5))
        if profile.get('headline'):
            document.add_paragraph(str(profile['headline']))
        contact = document.add_paragraph(' | '.join(str(profile[k]) for k in ('email', 'phone', 'location') if profile.get(k)))
        for link in profile.get('links', []):
            if contact.text:
                contact.add_run(' | ')
            _link(contact, link.get('label'), link.get('url', ''))
        if profile.get('summary'):
            document.add_paragraph(str(profile['summary']))
        for group in payload.get('sections', []):
            document.add_paragraph(str(group.get('title', '')), 'Heading 1')
            for entry in group.get('entries', []):
                label = ' · '.join(str(entry[k]) for k in ('title', 'organization', 'role') if entry.get(k))
                p = document.add_paragraph()
                p.paragraph_format.keep_with_next = bool(entry.get('summary') or entry.get('bullets'))
                p.add_run(label).bold = True
                dates = ' – '.join(str(entry[k]) for k in ('start', 'end') if entry.get(k))
                if dates:
                    p.add_run(' | ' + dates)
                if entry.get('url'):
                    p = document.add_paragraph(); _link(p, entry['url'], entry['url'])
                if entry.get('summary'):
                    document.add_paragraph(str(entry['summary']))
                for bullet in entry.get('bullets', []):
                    document.add_paragraph(str(bullet), 'List Bullet')
        _clean_metadata(document)
        available = _font_available(settings['body_font'])
        warnings = ['页数需在 Word 导出后检查；未按页数偏好缩小字号']
        if layout:
            warnings.append('PDF 模板仅重建提取到的布局特征；请逐页核对差异')
        if available is not True:
            warnings.append('指定字体在本机不可用或尚未确认，Word 可能替代字体；请核对导出结果：' + settings['body_font'])
        if settings.get('columns', 1) != 1:
            warnings.append('当前重建使用单栏正文，原模板多栏关系需人工核对')
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open('xb') as stream:
            document.save(stream)
        return {'ok': True, 'docx_path': str(output), 'path': str(output),
                'source_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
                'layout': settings, 'adjustments': adjustments,
                'font_validation': {'requested': requested_font, 'resolved': settings['body_font'], 'available': available},
                'warnings': warnings}
    except (ValueError, TypeError, KeyError, OSError) as error:
        return {'ok': False, 'error': str(error), 'docx_path': str(output)}
