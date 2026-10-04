"""Inspect and adapt plain DOCX templates; measure PDF layout without OCR.

Explicit mapping uses zero-based indexes from ``analyze_template``::

    {"paragraphs": {"0": "profile.name", "3": "section:work"},
     "tables": {"0:0:1": "profile.email"},
     "remove_paragraphs": [1], "remove_cells": ["0:1:1"],
     "preserve_paragraphs": [2], "preserve_cells": ["0:0:0"],
     "locations": {"header:0:0": "profile.contact"},
     "remove_locations": [], "preserve_locations": []}

The table key is table:row:column. Mapping a section replaces that location
with its complete, confirmed payload text. ``sections`` means all sections;
``profile.contact`` combines contacts. Preserve lists explicitly approve
static text. Unmapped nonempty text blocks output, including headers/footers.
Template text is never used as resume facts. Complex drawings/columns are
reported and require a different template or a separately reviewed rebuild.
"""
from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import re
import statistics

from docx import Document
from docx.oxml.ns import qn

PLACEHOLDER = re.compile(r'\{\{\s*([^{}]+?)\s*\}\}')
LABELS = {
    '姓名': 'name', 'name': 'name', '电话': 'phone', '手机': 'phone', 'phone': 'phone',
    'mobile': 'phone', '邮箱': 'email', '电子邮箱': 'email', 'email': 'email',
    '所在地': 'location', '地址': 'location', 'location': 'location',
    '求职方向': 'headline', 'headline': 'headline', '个人简介': 'summary',
    'summary': 'summary', '联系方式': 'contact', 'contact': 'contact',
}
HEADINGS = {
    '工作经历': 'work', '工作经验': 'work', 'work experience': 'work', 'experience': 'work',
    '教育经历': 'education', '教育背景': 'education', 'education': 'education',
    '项目经历': 'project', '项目经验': 'project', 'projects': 'project',
    '研究经历': 'research', 'research': 'research', 'teaching': 'teaching', '教学经历': 'teaching',
    '志愿经历': 'volunteer', 'volunteer': 'volunteer', '获奖经历': 'award', 'awards': 'award',
    '证书': 'certificate', 'certifications': 'certificate', 'publications': 'publication', '发表论文': 'publication',
}


def _normal(text):
    return text.strip().rstrip(':：').strip().lower()


def _text(paragraph):
    return ''.join('\n' if node.tag in (qn('w:br'), qn('w:cr')) else '\t' if node.tag == qn('w:tab') else node.text or ''
                   for node in paragraph._p.xpath('.//w:t | .//w:br | .//w:cr | .//w:tab'))


def _replace_span(paragraph, old, new):
    """Replace exact text across XML runs while keeping unaffected run styles."""
    nodes = paragraph._p.xpath('.//w:t')
    combined = ''.join(n.text or '' for n in nodes)
    starts = [m.start() for m in re.finditer(re.escape(old), combined)]
    for start in reversed(starts):
        end = start + len(old); offset = 0; first = None; last = None
        for index, node in enumerate(nodes):
            length = len(node.text or '')
            if first is None and offset <= start < offset + length:
                first = (index, start - offset)
            if offset < end <= offset + length:
                last = (index, end - offset); break
            offset += length
        if first is None or last is None:
            raise ValueError('无法定位文本修改范围')
        a, a_offset = first; b, b_offset = last
        prefix = (nodes[a].text or '')[:a_offset]; suffix = (nodes[b].text or '')[b_offset:]
        if a == b:
            nodes[a].text = prefix + new + suffix
        else:
            nodes[a].text = prefix + new
            for index in range(a + 1, b): nodes[index].text = ''
            nodes[b].text = suffix
        for index in range(a, b + 1): nodes[index].set(qn('xml:space'), 'preserve')
    return len(starts)


def _set_text(paragraph, value):
    from docx.text.run import Run
    _unlink(paragraph)
    runs = paragraph._p.xpath('./w:r')
    if runs:
        Run(runs[0], paragraph).text = value
        for run in runs[1:]: Run(run, paragraph).text = ''
    else: paragraph.add_run(value)


def _unlink(paragraph):
    # Remove inherited sample URL targets while preserving their formatted runs.
    for link in paragraph._p.xpath('.//w:hyperlink'):
        parent = link.getparent(); index = parent.index(link)
        for child in list(link): parent.insert(index, child); index += 1
        parent.remove(link)


def _normalize_breaks(paragraph):
    from docx.oxml import OxmlElement
    for node in list(paragraph._p.xpath('.//w:t')):
        if not re.search(r'[\n\t]', node.text or ''): continue
        parent = node.getparent(); index = parent.index(node)
        for piece in re.split(r'([\n\t])', node.text):
            if piece in ('\n', '\t'):
                element = OxmlElement('w:br' if piece == '\n' else 'w:tab')
            else:
                element = OxmlElement('w:t'); element.text = piece; element.set(qn('xml:space'), 'preserve')
            parent.insert(index, element); index += 1
        parent.remove(node)


def _link_urls(paragraph, urls):
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    while True:
        candidates = [n for n in paragraph._p.xpath('.//w:t') if n.getparent().getparent().tag != qn('w:hyperlink') and any(url in (n.text or '') for url in urls)]
        if not candidates: break
        node = candidates[0]
        text = node.text or ''
        pattern = '|'.join(re.escape(url) for url in sorted(urls, key=len, reverse=True) if url in text)
        if not pattern: continue
        run = node.getparent(); parent = run.getparent()
        if parent.tag == qn('w:hyperlink'): continue
        index = parent.index(run); cursor = 0
        fragments = []
        for match in re.finditer(pattern, text):
            if match.start() > cursor: fragments.append((text[cursor:match.start()], None))
            fragments.append((match.group(0), match.group(0))); cursor = match.end()
        if cursor < len(text): fragments.append((text[cursor:], None))
        # Other elements such as line breaks must stay in their original order.
        before = list(run)[:list(run).index(node)]; after = list(run)[list(run).index(node) + 1:]
        rpr = run.find(qn('w:rPr'))
        for prefix_nodes in (before,):
            content = [item for item in prefix_nodes if item.tag != qn('w:rPr')]
            if content:
                prefix_run = OxmlElement('w:r')
                if rpr is not None: prefix_run.append(deepcopy(rpr))
                for item in content: prefix_run.append(deepcopy(item))
                parent.insert(index, prefix_run); index += 1
        for fragment, url in fragments:
            fresh = OxmlElement('w:r')
            if rpr is not None: fresh.append(deepcopy(rpr))
            t = OxmlElement('w:t'); t.text = fragment; t.set(qn('xml:space'), 'preserve'); fresh.append(t)
            if url:
                link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)); link.append(fresh)
                parent.insert(index, link)
            else: parent.insert(index, fresh)
            index += 1
        if after:
            suffix_run = OxmlElement('w:r')
            if rpr is not None: suffix_run.append(deepcopy(rpr))
            for item in after: suffix_run.append(deepcopy(item))
            parent.insert(index, suffix_run)
        parent.remove(run)


def _clean_template_metadata(doc):
    for field in ('author', 'last_modified_by', 'title', 'subject', 'comments', 'keywords', 'category', 'identifier'):
        setattr(doc.core_properties, field, '')
    for part in list(doc.part.package.parts):
        for rid, rel in list(part.rels.items()):
            if rel.reltype.endswith('/customXml') or rel.reltype.endswith('/custom-properties') or rel.reltype.endswith('/comments'):
                del part.rels[rid]
            elif rel.reltype.endswith('/hyperlink') and hasattr(part, '_element') and rid not in part._element.xpath('.//@r:id'):
                del part.rels[rid]
        if hasattr(part, '_element'):
            for node in part._element.xpath('.//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference'):
                node.getparent().remove(node)


def _all_paragraphs(doc):
    """Yield unique body, nested-table, header and footer paragraphs."""
    from docx.text.paragraph import Paragraph
    seen = set()
    roots = [('body', doc._element.body, doc)]
    for index, section in enumerate(doc.sections):
        roots.extend((prefix, part._element, part) for prefix, part in _header_footer_parts(section, index))
    for prefix, root, parent in roots:
        for index, element in enumerate(root.xpath('.//w:p')):
            if element in seen: continue
            seen.add(element)
            yield f'{prefix}:{index}', Paragraph(element, parent)


def _header_footer_parts(section, index):
    """Enumerate definitions without triggering python-docx's lazy creation.

    Accessing an undefined Header/Footer ``_element`` creates a new part and
    sectPr reference. Only explicitly referenced variants may be accessed;
    linked variants inherit an earlier section's definition already visited.
    """
    parts = []
    variants = [('default', '', 'header', 'footer'),
                ('first', ':first', 'first_page_header', 'first_page_footer'),
                ('even', ':even', 'even_page_header', 'even_page_footer')]
    for kind in ('header', 'footer'):
        defined = {ref.get(qn('w:type'), 'default') for ref in section._sectPr.xpath(f'./w:{kind}Reference')}
        for variant, suffix, header_attribute, footer_attribute in variants:
            if variant in defined:
                attribute = header_attribute if kind == 'header' else footer_attribute
                parts.append((f'{kind}{suffix}:{index}', getattr(section, attribute)))
    return parts


def _safe_output(source, output):
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve(): raise ValueError('输出不能覆盖输入原件')
    if output.exists(): raise FileExistsError(f'输出已存在：{output.name}')
    if not source.is_file(): raise FileNotFoundError(str(source))
    output.parent.mkdir(parents=True, exist_ok=True)
    return source, output


def _save_new(doc, output):
    buffer = BytesIO(); doc.save(buffer)
    with Path(output).open('xb') as stream: stream.write(buffer.getvalue())


def _issues(doc):
    issues = []
    xml = doc._element.xml
    if 'txbxContent' in xml: issues.append('textboxes')
    if '<wp:anchor' in xml: issues.append('floating_objects')
    if any(int(cols.get(qn('w:num'), '1')) > 1 for cols in doc._element.xpath('.//w:cols')):
        issues.append('columns')
    if doc._element.xpath('.//w:tbl//w:tbl'): issues.append('nested_tables')
    if doc._element.xpath('.//w:vMerge') or doc._element.xpath('.//w:gridSpan'): issues.append('merged_cells')
    if doc._element.xpath('.//w:ins') or doc._element.xpath('.//w:del'): issues.append('tracked_changes')
    for index, section in enumerate(doc.sections):
        for _, part in _header_footer_parts(section, index):
            if 'txbxContent' in part._element.xml and 'textboxes' not in issues: issues.append('textboxes')
            if '<wp:anchor' in part._element.xml and 'floating_objects' not in issues: issues.append('floating_objects')
    return issues


def analyze_template(path):
    path = Path(path)
    if path.suffix.lower() == '.pdf': return pdf_layout_profile(path)
    if path.suffix.lower() != '.docx':
        return {'ok': False, 'format': path.suffix.lstrip('.'), 'issues': ['requires_conversion'], 'warnings': ['请先转换为 DOCX；原件保持不变。']}
    doc = Document(path); issues = _issues(doc)
    paragraphs = [{'index': i, 'text': _text(p), 'style': p.style.name,
                   'placeholders': PLACEHOLDER.findall(_text(p)),
                   'section': HEADINGS.get(_normal(_text(p)))} for i, p in enumerate(doc.paragraphs)]
    tables = [{'index': i, 'rows': [[{'key': f'{i}:{r}:{c}', 'text': cell.text} for c, cell in enumerate(row.cells)] for r, row in enumerate(table.rows)]} for i, table in enumerate(doc.tables)]
    return {'ok': not issues, 'format': 'docx', 'paragraphs': paragraphs, 'tables': tables,
            'locations': [{'location': location, 'text': _text(p)} for location, p in _all_paragraphs(doc)],
            'issues': issues, 'warnings': [f'复杂版式特征：{issue}；普通模板适配不能保证保真。' for issue in issues],
            'mapping_schema': 'paragraphs[index]=field; tables[table:row:column]=field; locations[body/header/footer:index]=field; remove/preserve_paragraphs; remove/preserve_cells; remove/preserve_locations',
            'allow_fact_import': False}


def _section_text(section, heading=False):
    lines = [section.get('title', '')] if heading else []
    for entry in section.get('entries', []):
        title = ' | '.join(str(entry.get(key, '')).strip() for key in ('title', 'organization', 'role') if entry.get(key))
        dates = ' – '.join(str(entry.get(key, '')) for key in ('start', 'end') if entry.get(key))
        lines.extend(value for value in [title, dates, entry.get('url', ''), entry.get('summary', '')] if value)
        lines.extend('• ' + str(bullet) for bullet in entry.get('bullets', []))
    return '\n'.join(lines)


def _value(field, payload):
    field = field.strip()
    if field == 'sections': return '\n\n'.join(_section_text(section, True) for section in payload.get('sections', []))
    if field.startswith('section:') or field.startswith('sections.'):
        section_id = field.split(':', 1)[1] if ':' in field else field.split('.', 1)[1]
        matches = [s for s in payload.get('sections', []) if s.get('id') == section_id]
        if len(matches) > 1: raise ValueError(f'栏目 id 重复：{section_id}')
        return _section_text(matches[0]) if matches else ''
    if field == 'profile.contact':
        profile = payload.get('profile', {})
        pieces = [str(profile[key]) for key in ('email', 'phone', 'location') if profile.get(key)]
        pieces.extend(str(item.get('url', '')) for item in profile.get('links', []) if item.get('url'))
        return ' | '.join(pieces)
    if field == 'profile.links':
        return ' | '.join((str(item.get('label', '')) + ': ' if item.get('label') else '') + str(item['url'])
                          for item in payload.get('profile', {}).get('links', []) if item.get('url'))
    parts = field.split('.')
    if len(parts) == 1 and parts[0] in ('name', 'headline', 'email', 'phone', 'location', 'summary'):
        parts.insert(0, 'profile')
    if len(parts) != 2 or parts[0] not in ('profile', 'target'):
        raise ValueError(f'不支持的字段映射：{field}')
    result = payload.get(parts[0], {}).get(parts[1], '')
    if isinstance(result, (dict, list)): raise ValueError(f'字段必须为文本：{field}')
    return str(result or '')


def adapt_docx(template_path, payload, output_path, mapping=None):
    source, output = _safe_output(template_path, output_path)
    report = analyze_template(source)
    if report['format'] != 'docx' or report['issues']:
        return dict(report, ok=False, output=None, mapping_required=False)
    doc = Document(source); mapping = mapping or {}; handled = set(); changes = []; unmapped = []
    paragraphs = list(doc.paragraphs)
    explicit = {str(k): v for k, v in mapping.get('paragraphs', {}).items()}
    remove = {str(k) for k in mapping.get('remove_paragraphs', [])}
    preserve = {str(k) for k in mapping.get('preserve_paragraphs', [])}
    cells_map = {str(k): v for k, v in mapping.get('tables', {}).items()}
    cells_remove = {str(k) for k in mapping.get('remove_cells', [])}
    cells_preserve = {str(k) for k in mapping.get('preserve_cells', [])}
    valid_p = {str(i) for i in range(len(paragraphs))}
    valid_cells = {f'{t}:{r}:{c}' for t, table in enumerate(doc.tables) for r, row in enumerate(table.rows) for c, cell in enumerate(row.cells)}
    if (set(explicit) | remove | preserve) - valid_p or (set(cells_map) | cells_remove | cells_preserve) - valid_cells:
        raise ValueError('映射位置不存在；请重新分析当前模板')

    def fill(p, field, location):
        old = _text(p); new = _value(field, payload); _set_text(p, new)
        handled.add(p._p); changes.append({'location': location, 'field': field, 'before': old, 'after': new})

    locations = dict(_all_paragraphs(doc)); mapped_locations = mapping.get('locations', {})
    removed_locations = set(mapping.get('remove_locations', [])); preserved_locations = set(mapping.get('preserve_locations', []))
    if (set(mapped_locations) | removed_locations | preserved_locations) - set(locations): raise ValueError('映射位置不存在；请重新分析当前模板')
    for location, p in locations.items():
        if location in mapped_locations: fill(p, mapped_locations[location], location)
        elif location in removed_locations:
            handled.add(p._p); _set_text(p, '')
        elif location in preserved_locations: handled.add(p._p)

    for index, p in enumerate(paragraphs):
        key = str(index)
        if key in explicit: fill(p, explicit[key], f'paragraph:{index}')
        elif key in remove:
            handled.add(p._p); p._p.getparent().remove(p._p)
        elif key in preserve: handled.add(p._p)
    # Recognized heading areas own every following ordinary paragraph up to
    # the next heading. Unknown headings are boundaries, never inferred facts.
    headings = [(i, HEADINGS.get(_normal(_text(p)))) for i, p in enumerate(paragraphs)
                if p._p not in handled and (HEADINGS.get(_normal(_text(p))) or p.style.name.lower().startswith('heading'))]
    for position, (index, section_id) in enumerate(headings):
        if not section_id: continue
        heading = paragraphs[index]; handled.add(heading._p)
        stop = headings[position + 1][0] if position + 1 < len(headings) else len(paragraphs)
        area = [p for p in paragraphs[index + 1:stop] if p._p not in handled and p._p.getparent() is not None]
        section_placeholder = any(field.strip() in (f'section:{section_id}', f'sections.{section_id}', 'sections') for p in area for field in PLACEHOLDER.findall(_text(p)))
        zone = [p for p in area if not PLACEHOLDER.search(_text(p))]
        value = _value(f'section:{section_id}', payload)
        if section_placeholder:
            for p in zone: handled.add(p._p); p._p.getparent().remove(p._p)
        elif value:
            if zone:
                fill(zone[0], f'section:{section_id}', f'paragraph:{index + 1}')
                for p in zone[1:]: handled.add(p._p); p._p.getparent().remove(p._p)
            else:
                from docx.text.paragraph import Paragraph
                element = deepcopy(heading._p); heading._p.addnext(element)
                p = Paragraph(element, doc); p.style = doc.styles['Normal']; _set_text(p, value); handled.add(p._p)
        else:
            for p in [heading] + zone: handled.add(p._p); p._p.getparent().remove(p._p)
    for t, table in enumerate(doc.tables):
        seen_cells = set()
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if cell._tc in seen_cells: continue
                seen_cells.add(cell._tc); key = f'{t}:{r}:{c}'
                field = cells_map.get(key)
                if field is None and c == 1 and len(row.cells) == 2:
                    label = _normal(row.cells[0].text)
                    if label in LABELS:
                        field = 'profile.' + LABELS[label]
                        handled.update(p._p for p in row.cells[0].paragraphs)
                    elif label in HEADINGS:
                        field = 'section:' + HEADINGS[label]
                        handled.update(p._p for p in row.cells[0].paragraphs)
                if field is not None:
                    fill(cell.paragraphs[0], field, f'table:{key}')
                    for p in cell.paragraphs[1:]: _set_text(p, ''); handled.add(p._p)
                elif key in cells_remove:
                    for p in cell.paragraphs: _set_text(p, ''); handled.add(p._p)
                elif key in cells_preserve: handled.update(p._p for p in cell.paragraphs)
    for location, p in _all_paragraphs(doc):
        if p._p in handled or not _text(p).strip(): continue
        original = _text(p); matches = list(PLACEHOLDER.finditer(original))
        if matches:
            # A paragraph mixing a placeholder and arbitrary sample text is
            # ambiguous. Recognized field prefixes such as "姓名：" are safe.
            remainder = PLACEHOLDER.sub('', original).strip(' \t\n|,，;；:：-–')
            if remainder and _normal(remainder) not in LABELS:
                unmapped.append({'location': location, 'text': original, 'reason': 'mixed_sample_text'}); continue
            try:
                _unlink(p)
                for match in reversed(matches): _replace_span(p, match.group(0), _value(match.group(1), payload))
            except ValueError as exc:
                unmapped.append({'location': location, 'text': original, 'reason': str(exc)}); continue
            handled.add(p._p); changes.append({'location': location, 'before': original, 'after': _text(p)})
        elif re.match(r'^([^:：]+)([:：])', original.strip()) and _normal(re.match(r'^([^:：]+)([:：])', original.strip()).group(1)) in LABELS:
            match = re.match(r'^([^:：]+)([:：])', original.strip())
            field = 'profile.' + LABELS[_normal(match.group(1))]
            new = match.group(1) + match.group(2) + _value(field, payload)
            _set_text(p, new); handled.add(p._p)
            changes.append({'location': location, 'field': field, 'before': original, 'after': new})
        elif _normal(original) in ('简历', '个人简历', 'resume', 'curriculum vitae'):
            handled.add(p._p)
        else: unmapped.append({'location': location, 'text': original, 'reason': 'unmapped_text'})
    expected_urls = {str(item['url']) for item in payload.get('profile', {}).get('links', []) if item.get('url')}
    expected_urls.update(str(entry['url']) for section in payload.get('sections', []) for entry in section.get('entries', []) if entry.get('url'))
    rendered = '\n'.join(_text(p) for _, p in _all_paragraphs(doc))
    for url in sorted(expected_urls):
        if url not in rendered: unmapped.append({'location': 'payload.links', 'text': url, 'reason': 'missing_link_slot'})
    if unmapped:
        return {'ok': False, 'mapping_required': True, 'unmapped': unmapped, 'warnings': ['存在未映射文字；请明确替换、删除或保留位置后重试。'], 'output': None, 'changes': changes}
    for _, p in _all_paragraphs(doc):
        if p._p in handled:
            _normalize_breaks(p); _link_urls(p, expected_urls)
    _clean_template_metadata(doc)
    _save_new(doc, output)
    return {'ok': True, 'mapping_required': False, 'unmapped': [], 'warnings': ['普通段落栏目内容沿用首个样例段落的格式；复杂布局未重建。'],
            'output': str(output), 'changes': changes, 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest()}


def _color_hex(color):
    if color is None: return '000000'
    if not isinstance(color, (tuple, list)): color = [color]
    values = [float(v) for v in color]
    if len(values) == 1: values *= 3
    elif len(values) == 4:
        c, m, y, k = values; values = [(1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)]
    return ''.join(f'{round(max(0, min(1, v)) * 255):02X}' for v in values[:3])


def pdf_layout_profile(path):
    import pdfplumber
    pages = []; all_chars = []; headings = []; section_heading_chars = []; issues = []; warnings = []; scanned = []; adjustments = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages):
            chars = [char for char in page.chars if str(char.get('text', '')).strip()]
            all_chars.extend(chars)
            bounds = None
            if chars:
                bounds = {'left': min(c['x0'] for c in chars), 'right': max(c['x1'] for c in chars),
                          'top': min(c['top'] for c in chars), 'bottom': max(c['bottom'] for c in chars)}
            else: scanned.append(index + 1)
            pages.append({'page': index + 1, 'width_pt': float(page.width), 'height_pt': float(page.height),
                          'text': page.extract_text() or '', 'bounds_pt': bounds, 'character_count': len(chars),
                          'images': len(page.images), 'vector_shapes': len(page.rects) + len(page.curves) + len(page.lines)})
            for line in page.extract_text_lines(return_chars=True) if chars else []:
                section_id = HEADINGS.get(_normal(line['text']))
                if not section_id: continue
                line_chars = [c for c in line.get('chars', []) if str(c.get('text', '')).strip()]
                section_heading_chars.extend(line_chars)
                heading = {'page': index + 1, 'text': line['text'], 'section': section_id}
                if line_chars:
                    heading.update({'size_pt': Counter(round(float(c['size']), 2) for c in line_chars).most_common(1)[0][0],
                                    'color': Counter(_color_hex(c.get('non_stroking_color')) for c in line_chars).most_common(1)[0][0],
                                    'font': Counter(c['fontname'] for c in line_chars).most_common(1)[0][0]})
                headings.append(heading)
            words = page.extract_words() if chars else []
            # Two separated text bands sharing several vertical lines are a
            # multi-column signal. This is a heuristic, reported as such.
            bands = {}
            for word in words:
                key = round(word['top'] / 4)
                bands.setdefault(key, []).append(word)
            split_rows = 0
            for line in bands.values():
                line.sort(key=lambda w: w['x0'])
                if any(b['x0'] - a['x1'] > page.width * .16 for a, b in zip(line, line[1:])): split_rows += 1
            if split_rows >= 3 and 'columns' not in issues: issues.append('columns')
    if not pages: raise ValueError('PDF 无页面')
    first = pages[0]; scale = 2.54 / 72
    layout = {'paper_width_cm': round(first['width_pt'] * scale, 4), 'paper_height_cm': round(first['height_pt'] * scale, 4)}
    layout['text_bounds_pt'] = [{'page': page['page'], 'bounds_pt': page['bounds_pt']} for page in pages]
    if len({(p['width_pt'], p['height_pt']) for p in pages}) > 1: issues.append('mixed_page_sizes')
    if scanned:
        issues.append('ocr_required'); warnings.append('以下页面无可读取文字层，需要 OCR 或人工转录：' + ', '.join(map(str, scanned)))
    if all_chars:
        sizes = Counter(round(float(c['size']), 2) for c in all_chars)
        body_size = sizes.most_common(1)[0][0]
        body_chars = [c for c in all_chars if abs(float(c['size']) - body_size) < .1]
        font = Counter(c['fontname'] for c in body_chars).most_common(1)[0][0]
        large = [c for c in all_chars if float(c['size']) >= body_size * 1.15]
        # A name can be larger than every section heading. Recognized section
        # lines supply semantic evidence, so use their actual glyph styles
        # before falling back to the existing larger-font heuristic.
        heading_candidates = section_heading_chars or large
        heading_size = Counter(round(float(c['size']), 2) for c in heading_candidates).most_common(1)[0][0] if heading_candidates else body_size
        heading_chars = [c for c in heading_candidates if abs(float(c['size']) - heading_size) < .1]
        heading_color = Counter(_color_hex(c.get('non_stroking_color')) for c in heading_chars).most_common(1)[0][0] if heading_chars else '000000'
        margins = {'top': [], 'bottom': [], 'left': [], 'right': []}
        for page in pages:
            b = page['bounds_pt']
            if b:
                margins['top'].append(b['top']); margins['bottom'].append(page['height_pt'] - b['bottom'])
                margins['left'].append(b['left']); margins['right'].append(page['width_pt'] - b['right'])
                page['measured_content_gaps_cm'] = {'top': round(max(0, b['top']) * scale, 4),
                                                    'bottom': round(max(0, page['height_pt'] - b['bottom']) * scale, 4),
                                                    'left': round(max(0, b['left']) * scale, 4),
                                                    'right': round(max(0, page['width_pt'] - b['right']) * scale, 4)}
        measured_gaps = {key: round(max(0, statistics.median(values)) * scale, 4) for key, values in margins.items()}
        inferred_margins = dict(measured_gaps)
        # Whitespace after a short paragraph is not evidence of a physical
        # page margin. Only plain single-column pages with plausible leading
        # margins receive a conservative, explicitly reported symmetric guess.
        plain_single_column = not issues and not scanned and all(page['images'] == 0 and page['vector_shapes'] == 0 for page in pages)
        if plain_single_column:
            for leading, trailing, dimension in [('left', 'right', 'paper_width_cm'), ('top', 'bottom', 'paper_height_cm')]:
                near, far = measured_gaps[leading], measured_gaps[trailing]
                if .5 <= near <= 3.5 and far >= layout[dimension] * .25 and far >= near * 2.5 and far - near >= 2:
                    inferred_margins[trailing] = near
                    reason = '单栏且无图片/矢量图；末端留白达到页面尺寸25%、前端留白2.5倍且多出至少2cm，推断未填内容空白，使用对称边距回退。'
                    adjustments.append({'field': f'margins_cm.{trailing}', 'from': far, 'to': near, 'reason': reason})
                    warnings.append(f'{trailing} 测量留白 {far}cm 可能包含未填内容，推断采用与 {leading} 对称的 {near}cm 边距；需人工确认，原始文字边界与留白已保留。')
        if not adjustments:
            warnings.append('物理页边距无法由文字边界唯一确定；未满足保守对称回退条件，保留测量留白，页边距推断需人工确认。')
        layout.update({'margins_cm': inferred_margins, 'measured_content_gaps_cm': measured_gaps,
                       'body_font': re.sub(r'^[A-Z]{6}\+', '', font), 'body_size_pt': body_size,
                       'heading_size_pt': heading_size, 'heading_color': heading_color,
                       'columns': 2 if 'columns' in issues else 1})
        warnings.append('PDF 以文字边界估计页边距、以主要字形估计字体；重建不保留绝对定位、图片和矢量装饰，字体需本机可用。')
    layout['issues'] = list(issues)
    layout['adjustments'] = adjustments
    if 'columns' in issues: warnings.append('检测到可能的多栏排版；需人工确认，当前重建器不能保证保留多栏关系。')
    return {'ok': bool(all_chars) and not scanned, 'format': 'pdf', 'requires_ocr': bool(scanned),
            'ocr_pages': scanned, 'layout': layout, 'pages': pages, 'headings': headings,
            'issues': issues, 'warnings': warnings, 'adjustments': adjustments, 'allow_fact_import': False}
