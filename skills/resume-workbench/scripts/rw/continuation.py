"""Continue a latest Word document independently of the resume fact history.

replacements is {old: new} (each old must occur exactly once), or a list of
{"old": "...", "new": "...", "location": "body:3", "count": 1} records.
Location comes from inspect_input.structure.paragraphs. Explicit count allows
the caller to approve repeated occurrences. All edits are validated together
before writing; overlapping edits are rejected. Newly edited text is pending.
"""
from pathlib import Path
import hashlib
from docx import Document
from .templates import _safe_output, _save_new, _all_paragraphs, _text, _normalize_breaks


def continue_docx(latest_path, output_path, replacements=None):
    source, output = _safe_output(latest_path, output_path)
    if source.suffix.lower() != '.docx': raise ValueError('续编需要最新版 DOCX 文件')
    doc = Document(source); paragraphs = list(_all_paragraphs(doc))
    original = {location: _text(p) for location, p in paragraphs}
    if replacements is None: records = []
    elif isinstance(replacements, dict): records = [{'old': old, 'new': new} for old, new in replacements.items()]
    elif isinstance(replacements, list): records = replacements
    else: raise ValueError('replacements 必须为字典或修改列表')
    plans = []; occupied = {}
    for record in records:
        old, new = record.get('old'), record.get('new')
        if not isinstance(old, str) or not old or not isinstance(new, str): raise ValueError('修改需要非空 old 和文本 new')
        location_filter = record.get('location'); expected = record.get('count', 1)
        if not isinstance(expected, int) or isinstance(expected, bool) or expected < 1: raise ValueError('count 必须为正整数')
        selected = [(loc, p) for loc, p in paragraphs if location_filter is None or loc == location_filter]
        matches = [(loc, p, original[loc].count(old)) for loc, p in selected if old in original[loc]]
        actual = sum(count for _, _, count in matches)
        if actual != expected: raise ValueError(f'明确匹配失败：期望 {expected} 次，实际 {actual} 次；请指定准确原文/位置/count。')
        for location, p, count in matches:
            start = 0
            while True:
                start = original[location].find(old, start)
                if start < 0: break
                end = start + len(old)
                if any(start < right and end > left for left, right in occupied.get(location, [])):
                    raise ValueError('修改范围重叠；请合并为一次明确修改')
                occupied.setdefault(location, []).append((start, end)); start = end
            plans.append((location, p, old, new, count))
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    if not plans:
        with output.open('xb') as stream: stream.write(source.read_bytes())
    else:
        # Apply each paragraph from right to left so text inserted by an edit
        # cannot become the target of a different edit in this same request.
        for location, p in paragraphs:
            local = [plan for plan in plans if plan[0] == location]
            if not local: continue
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            # Account for semantic breaks/tabs when calculating text offsets.
            for node in list(p._p.xpath('.//w:br | .//w:cr | .//w:tab')):
                text_node = OxmlElement('w:t'); text_node.text = '\t' if node.tag == qn('w:tab') else '\n'
                node.getparent().replace(node, text_node)
            spans = []
            for _, _, old, new, _ in local:
                start = 0
                while True:
                    start = original[location].find(old, start)
                    if start < 0: break
                    spans.append((start, old, new)); start += len(old)
            # Replace using an exact offset helper to avoid replacing newly
            # introduced matching substrings elsewhere in the paragraph.
            for start, old, new in sorted(spans, reverse=True):
                nodes = p._p.xpath('.//w:t'); cursor = 0; begin = finish = None; end = start + len(old)
                for i, node in enumerate(nodes):
                    length = len(node.text or '')
                    if begin is None and cursor <= start < cursor + length: begin = (i, start - cursor)
                    if cursor < end <= cursor + length: finish = (i, end - cursor); break
                    cursor += length
                a, x = begin; b, y = finish
                prefix = (nodes[a].text or '')[:x]; suffix = (nodes[b].text or '')[y:]
                nodes[a].text = prefix + new + (suffix if a == b else '')
                if a != b:
                    for i in range(a + 1, b): nodes[i].text = ''
                    nodes[b].text = suffix
                for i in range(a, b + 1): nodes[i].set(qn('xml:space'), 'preserve')
            _normalize_breaks(p)
        _save_new(doc, output)
    reread = {location: _text(p) for location, p in _all_paragraphs(Document(output))}
    changes = [{'location': location, 'before': original[location], 'after': reread.get(location, '')}
               for location in original if original[location] != reread.get(location, '')]
    return {'ok': True, 'output': str(output), 'source_sha256': source_hash,
            'output_sha256': hashlib.sha256(output.read_bytes()).hexdigest(), 'changes': changes,
            'pending_facts': [{'text': change['after'], 'status': 'pending', 'location': change['location'], 'source': str(output)} for change in changes],
            'warnings': ['续编以最新版 Word 为准；改动回读为待确认资料，不自动更新或确认事实。'] if changes else [],
            'history_required': False, 'text': '\n'.join(reread.values())}
