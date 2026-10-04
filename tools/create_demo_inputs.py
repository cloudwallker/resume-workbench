"""Generate clearly fictional test inputs; never use real candidate data."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from pypdf import PdfWriter

ROOT = Path(__file__).resolve().parents[1]
EXAMPLES = ROOT / 'skills/resume-workbench/assets/examples'
FIXTURES = ROOT / 'tests/fixtures'


def bullet(identifier, text, **extra):
    return {'id': identifier, 'text': text, 'status': 'confirmed', 'source_ids': ['user-1'], 'contribution': 'individual', **extra}


def experience(identifier, kind, title, organization, role, start, end, bullets, tags=None):
    return {'id': identifier, 'kind': kind, 'title': title, 'organization': organization, 'role': role,
            'start': start, 'end': end, 'status': 'confirmed', 'source_ids': ['user-1'], 'tags': tags or [], 'bullets': bullets}


def master(name, headline, language, seniority, experiences):
    return {'schema_version': '1.0', 'profile': {'name': name, 'headline': headline,
            'email': 'candidate@example.com', 'location': '示例城市' if language == 'zh' else 'Example City',
            'summary': '专注清晰交付、协作与流程改进。' if language == 'zh' else 'Operations professional focused on clear delivery, collaboration and process improvement.'},
            'target': {'role': headline, 'language': language, 'market': 'CN' if language == 'zh' else 'US',
                       'seniority': seniority, 'paper': 'A4' if language == 'zh' else 'Letter', 'page_limit': 1},
            'sources': [{'id': 'user-1', 'type': 'user', 'reference': '虚构验收访谈 / fictional acceptance interview',
                         'note': '全部信息为流程演示虚构资料，不属于真实求职者。'}],
            'experiences': experiences}


def create():
    EXAMPLES.mkdir(parents=True, exist_ok=True)
    FIXTURES.mkdir(parents=True, exist_ok=True)
    operations = master('示例求职者（虚构）', '活动运营专员', 'zh', 'experienced', [
        experience('work-1', 'work', '活动运营与复盘', '示例服务机构（虚构）', '运营专员', '2023-06', '2025-08', [
            bullet('w1-b1', '整理活动报名资料并交付周报，供团队复盘使用。'),
            bullet('w1-b2', '协调场地与物料交接，按活动清单检查执行进度。'),
            bullet('w1-b3', '维护常见问题说明，协助团队统一参与者答复口径。', contribution='assisted')], ['运营', '活动', '沟通', '复盘']),
        experience('education-1', 'education', '管理学学士', '示例大学（虚构）', '工商管理', '2019-09', '2023-06', [], ['管理']),
        experience('volunteer-1', 'volunteer', '社区阅读活动', '示例社区（虚构）', '志愿者', '2022-03', '2022-10', [
            bullet('v1-b1', '协助登记图书并整理借阅记录，向活动负责人交付台账。', contribution='assisted')], ['活动'])])
    technical = master('示例学生（虚构）', '软件开发实习生', 'zh', 'student', [
        experience('education-1', 'education', '计算机科学与技术本科', '示例大学（虚构）', '本科在读', '2023-09', 'Present', [], ['计算机']),
        experience('project-1', 'project', '课程任务管理项目', '课程团队（虚构）', '后端开发成员', '2025-03', '2025-06', [
            bullet('p1-b1', '实现任务状态校验与列表查询接口，编写接口使用说明。'),
            bullet('p1-b2', '为边界输入添加单元测试，协助团队完成课程演示。', contribution='assisted')], ['Python', '接口', '测试']),
        experience('project-2', 'project', '校园资料整理脚本', '个人学习项目（虚构）', '开发者', '2025-07', '2025-08', [
            bullet('p2-b1', '编写文件分类与重名处理脚本，保留操作日志供核对。')], ['Python', '文件'])])
    english = master('Alex Example (Fictional)', 'Operations Coordinator', 'en', 'experienced', [
        experience('work-1', 'work', 'Event Operations', 'Example Services (Fictional)', 'Operations Coordinator', '2021-06', '2025-08', [
            bullet('w1-b1', 'Coordinated venue and materials handovers using a shared event checklist.'),
            bullet('w1-b2', 'Compiled weekly reports for the team to review registrations and unresolved issues.'),
            bullet('w1-b3', 'Maintained response guidance and supported consistent participant communication.', contribution='assisted')], ['operations', 'events', 'communication']),
        experience('education-1', 'education', 'BA in Business Administration', 'Example University (Fictional)', 'Graduate', '2017-09', '2021-06', [], ['business'])])
    for filename, data in [('operations-zh.json', operations), ('technical-zh.json', technical), ('experienced-en.json', english)]:
        (EXAMPLES / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (EXAMPLES / 'jd-operations.txt').write_text('岗位：活动运营专员\n职责：活动执行、报名资料整理、沟通协调、周报及复盘。\n必要要求：清晰书面表达和活动协调经历。\n优先条件：CRM 系统使用经验。\n', encoding='utf-8')
    (FIXTURES / 'operations-description.txt').write_text('以下全部为虚构验收资料。示例求职者，目标活动运营专员，邮箱 candidate@example.com。2023年6月至2025年8月在示例服务机构做运营专员，整理活动报名资料和周报；协调场地物料；协助团队维护答复口径。2019年9月至2023年6月在示例大学读工商管理本科。2022年3月至10月协助社区阅读活动登记图书和借阅记录。不知道可量化指标，不写数字成果。', encoding='utf-8')
    doc = Document()
    doc.styles['Normal'].font.name = 'Microsoft YaHei'; doc.styles['Normal'].font.size = Pt(11)
    doc.add_paragraph('姓名：样例姓名')
    doc.add_paragraph('联系方式：样例邮箱')
    doc.add_paragraph('求职方向：样例岗位')
    doc.add_paragraph('所在地：样例城市')
    doc.add_paragraph('个人简介：样例简介')
    doc.add_heading('工作经历', level=1)
    doc.add_paragraph('样例公司的样例职责，必须被全部替换。')
    doc.add_heading('教育经历', level=1)
    doc.add_paragraph('样例学校')
    doc.add_heading('志愿经历', level=1)
    doc.add_paragraph('样例志愿活动')
    doc.save(FIXTURES / 'paragraph-template.docx')
    table_doc = Document()
    table_doc.styles['Normal'].font.name = 'Microsoft YaHei'; table_doc.styles['Normal'].font.size = Pt(11)
    table_doc.add_heading('{{profile.name}}', level=0)
    table_doc.add_paragraph('{{profile.headline}}')
    table_doc.add_paragraph('{{profile.contact}}')
    table_doc.add_paragraph('{{profile.location}}')
    table_doc.add_paragraph('{{profile.summary}}')
    table = table_doc.add_table(rows=3, cols=2); table.style = 'Light Shading Accent 1'
    for row, kind, label in [(0, 'work', '工作经历'), (1, 'education', '教育经历'), (2, 'volunteer', '志愿经历')]:
        table.cell(row, 0).text = label
        table.cell(row, 1).text = '{{section:' + kind + '}}'
    table_doc.save(FIXTURES / 'table-template.docx')
    latest = Document(); p = latest.add_paragraph(); p.add_run('原来的手工文案').bold = True
    p.add_run('；保留用户手工修改的颜色。').font.color.rgb = RGBColor.from_string('245B78')
    latest.save(FIXTURES / 'manually-edited.docx')
    (FIXTURES / 'changes.json').write_text(json.dumps([{'old': '原来的手工文案', 'new': '用户指定的新文案'}], ensure_ascii=False), encoding='utf-8')
    writer = PdfWriter(); writer.add_blank_page(width=595, height=842)
    with (FIXTURES / 'no-text-template.pdf').open('wb') as stream: writer.write(stream)
    return {'examples': str(EXAMPLES), 'fixtures': str(FIXTURES)}


if __name__ == '__main__':
    print(json.dumps(create(), ensure_ascii=False))
